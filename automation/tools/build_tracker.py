#!/usr/bin/env python3
"""Add one alert's screening results to the master tracker workbook and render the email table.

  build_tracker.py <rows.json> <tracker.xlsx> --sheet <tab name> [--html <out.html> --title <text>]

The workbook follows the Google Sheet "Inovi – MyBidMatch Initial Solicitation Screening
Tracker": one tab per alert (22 columns) plus a single "Column Guide" tab. The tab <tab name>
is created as the first tab, replacing a tab of the same name (a re-run of the same alert).
Other tabs, including the reviewer's "Final Decision / Next Action" entries, are kept.

rows.json is a JSON list of objects keyed by the column headers below. "Sl. No." is filled
automatically; "Final Decision / Next Action" is always left blank (the reviewer fills it).
An optional "Manual Review Note" key per row feeds the email's "Needs manual review" list.
--html writes an email-ready HTML summary table (a subset of the columns) grouped by verdict.
Exits non-zero if a row is missing a column or has an invalid enum value.
"""
import argparse
import html
import json
import os
import re
import sys
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

GUIDE = [
    ("Sl. No.", "Sequential reference number."),
    ("Review Date", "Date the solicitation is screened."),
    ("Solicitation Title", "Official solicitation title."),
    ("Solicitation Number", "Official RFP/RFQ or reference number."),
    ("Solicitation Type", "RFP, RFQ, IFB, RFI, etc."),
    ("Agency", "Name of the issuing agency."),
    ("Scope Summary", "Brief description of the required solution, services and objectives."),
    ("Go Reasons / Strengths", "Key reasons supporting pursuit of the opportunity."),
    ("No-Go Reasons / Weaknesses", "Major barriers, risks or reasons not to pursue."),
    ("Mandatory Requirements Summary", "Key eligibility, experience, certification, insurance and technical requirements."),
    ("Salesforce Fit", "Whether Salesforce products and technology can meet the technical requirements."),
    ("Inovi Fit", "Whether Inovi can deliver the solution using its actual capabilities, employee qualifications and past performance."),
    ("Source", "MyBidMatch, GovWin, SAM.gov or another source."),
    ("Source Link", "Direct link to the solicitation or official listing."),
    ("Solicitation Issue Date", "Official publication date."),
    ("Submission Deadline", "Submission date, time and time zone, where available."),
    ("Capability Fit", "Overall alignment with Inovi's capabilities, resources and business focus."),
    ("Potential Disqualifiers / Gaps", "Mandatory eligibility gaps, missing qualifications or significant risks."),
    ("Evidence (Section + PDF Page)", "Supporting solicitation section and PDF page for key findings."),
    ("Mandatory Requirements Status", "Met, Not Met or Unclear, based on available evidence."),
    ("Screening Verdict", "Automated recommendation: Pursue, Maybe or Reject."),
    ("Final Decision / Next Action", "Human decision and concise next step (filled by the reviewer, not the flow)."),
]
RULES = [
    "Rules for every entry:",
    "• Keep each summary or description to a maximum of 5 lines.",
    "• Use concise, factual information supported by the solicitation and the four approved Inovi documents.",
    '• If information is not stated in the solicitation, enter exactly: "Not available in the solicitation."',
    '• If information cannot be verified, enter: "Needs verification."',
    "• Include evidence references wherever available; never invent facts, qualifications or page numbers.",
    "• Keep the automated Screening Verdict separate from the human Final Decision / Next Action.",
]
HEADERS = [h for h, _ in GUIDE]
ENUMS = {"Mandatory Requirements Status": {"Met", "Not Met", "Unclear"},
         "Screening Verdict": {"Pursue", "Maybe", "Reject"}}
VERDICT_FILL = {"Pursue": "C6EFCE", "Maybe": "FFEB9C", "Reject": "FFC7CE"}
WIDTHS = {"Sl. No.": 7, "Review Date": 14, "Solicitation Title": 34, "Solicitation Number": 18,
          "Solicitation Type": 14, "Agency": 22, "Source": 12, "Source Link": 30,
          "Solicitation Issue Date": 14, "Submission Deadline": 18,
          "Mandatory Requirements Status": 14, "Screening Verdict": 12}

def solid(rgb):
    """Opaque fill: openpyxl stores 6-digit colours with alpha 00, which some Excel builds show as no fill."""
    return PatternFill("solid", fgColor="FF" + rgb, bgColor="FF" + rgb)


HEAD_FONT = Font(bold=True, color="FF000000")
HEAD_FILL = solid("DDEBF7")
THIN = Side(style="thin", color="FFBFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


MAX_CELL = 32000  # Excel's limit is 32,767 characters


def validate(rows):
    if not isinstance(rows, list) or not rows:
        return ["rows.json must be a non-empty JSON list of objects"]
    problems = []
    for i, r in enumerate(rows, 1):
        if not isinstance(r, dict):
            problems.append(f"row {i}: not an object")
            continue
        for h in HEADERS[1:-1]:
            v = r.get(h)
            if not isinstance(v, str) or not v.strip():
                problems.append(f"row {i}: missing or non-text '{h}'")
            elif len(str(v)) > MAX_CELL:
                problems.append(f"row {i}: '{h}' is longer than {MAX_CELL} characters")
        for h, allowed in ENUMS.items():
            if r.get(h) and r[h] not in allowed:
                problems.append(f"row {i}: '{h}' = {r[h]!r}, expected one of {sorted(allowed)}")
        note = r.get("Manual Review Note")
        if note is not None and not isinstance(note, str):
            problems.append(f"row {i}: 'Manual Review Note' must be text")
        if r.get("Mandatory Requirements Status") == "Not Met" and r.get("Screening Verdict") == "Pursue":
            problems.append(f"row {i}: 'Not Met' mandatory requirements cannot have verdict 'Pursue'")
    return problems


def as_text(cell):
    """Keep untrusted text from being interpreted as a spreadsheet formula."""
    if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"


def fill_tracker(ws, rows):
    ws.append(HEADERS)
    for c in ws[1]:
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, WRAP, BOX
    for n, r in enumerate(rows, 1):
        values = [n] + [r.get(h, "") for h in HEADERS[1:-1]] + [""]
        ws.append(values)
        for c in ws[ws.max_row]:
            c.alignment, c.border = WRAP, BOX
            as_text(c)
        verdict = ws.cell(row=ws.max_row, column=HEADERS.index("Screening Verdict") + 1)
        if verdict.value in VERDICT_FILL:
            verdict.fill = solid(VERDICT_FILL[verdict.value])
            verdict.font = Font(bold=True)
        link = ws.cell(row=ws.max_row, column=HEADERS.index("Source Link") + 1)
        if str(link.value).startswith("http"):
            link.hyperlink = str(link.value)
            link.font = Font(color="FF0563C1", underline="single")
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = WIDTHS.get(h, 40)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions


def fill_guide(g):
    g.append(["No.", "Column Header", "What to Record"])
    for c in g[1]:
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
    for i, (h, desc) in enumerate(GUIDE, 1):
        g.append([i, h, desc])
    g.append([])
    for line in RULES:
        g.append(["", line, ""])
    g.column_dimensions["A"].width = 6
    g.column_dimensions["B"].width = 34
    g.column_dimensions["C"].width = 100
    for row in g.iter_rows():
        for c in row:
            c.alignment = WRAP


class WorkbookOpen(Exception):
    pass


def add_apply_flags(path):
    """openpyxl omits applyFont/applyFill/applyBorder on cell formats; some desktop Excel builds
    then draw no fill (Google Sheets ignores the flags). Add them to xl/styles.xml in place."""
    def flag(m):
        xf = m.group(0)
        for attr, key in (("applyFont", "fontId"), ("applyFill", "fillId"), ("applyBorder", "borderId")):
            if attr not in xf and not re.search(r'\b%s="0"' % key, xf):
                xf = xf.replace("<xf ", '<xf %s="1" ' % attr, 1)
        return xf

    def fix(styles):
        return re.sub(r"<cellXfs\b.*?</cellXfs>",
                      lambda m: re.sub(r"<xf\b[^>]*>", flag, m.group(0)), styles, flags=re.S)

    tmp = path + ".styles"
    try:
        with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "xl/styles.xml":
                    data = fix(data.decode("utf-8")).encode("utf-8")
                dst.writestr(item, data)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def save_workbook(rows, path, sheet):
    folder, name = os.path.split(os.path.abspath(path))
    if os.path.exists(os.path.join(folder, "~$" + name)):
        # Excel's owner file: saving now would be overwritten (and the new tab lost) when Excel saves.
        raise WorkbookOpen(f"{name} is open in Excel; close it and re-run")
    if os.path.exists(path):
        wb = load_workbook(path, rich_text=True)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for existing in wb.sheetnames:
        if existing.lower() == sheet.lower():  # Excel tab names are case-insensitive
            del wb[existing]
    fill_tracker(wb.create_sheet(sheet, 0), rows)
    if "Column Guide" not in wb.sheetnames:
        fill_guide(wb.create_sheet("Column Guide"))
    wb.active = 0
    # Write to a temp file, then rename, so a crash never leaves a half-written master workbook.
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(path)))
    os.close(fd)
    try:
        wb.save(tmp)
        add_apply_flags(tmp)
        os.chmod(tmp, 0o644)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):  # only left behind when a step above failed
            os.remove(tmp)


EMAIL_COLS = ["Solicitation Title", "Solicitation Number", "Agency", "Submission Deadline",
              "Mandatory Requirements Status", "Key Reason", "Source Link"]


def render_html(rows, title, tracker_name, sheet):
    e = lambda v: html.escape(str(v))
    counts = {v: sum(r["Screening Verdict"] == v for r in rows) for v in ("Pursue", "Maybe", "Reject")}
    out = [f"<h2>{e(title)}</h2>",
           f"<p><b>Solicitations reviewed: {len(rows)}</b> (Pursue {counts['Pursue']}, "
           f"Maybe {counts['Maybe']}, Reject {counts['Reject']})</p>"]
    for verdict in ("Pursue", "Maybe", "Reject"):
        group = [r for r in rows if r["Screening Verdict"] == verdict]
        if not group:
            continue
        out.append(f'<h3 style="background:#{VERDICT_FILL[verdict]};padding:4px 8px">{verdict} ({len(group)})</h3>')
        # Plain table attributes keep the HTML short: the daily agent passes it to Gmail as text.
        out.append('<table border="1" cellpadding="5" cellspacing="0" style="border-collapse:collapse;font-size:13px">')
        out.append('<tr bgcolor="#d9e1f2">' + "".join(f"<th align=left>{c}</th>" for c in EMAIL_COLS) + "</tr>")
        for r in group:
            reason = r["No-Go Reasons / Weaknesses"] if verdict == "Reject" else r["Go Reasons / Strengths"]
            reason = reason if len(reason) <= 220 else reason[:217].rsplit(" ", 1)[0] + "…"  # full text is in the workbook
            status = r["Mandatory Requirements Status"]
            status_html = e(status) if status == "Met" else f'<b style="color:#c00000">ATTENTION: {e(status)}</b>'
            link = r["Source Link"]
            link_html = f'<a href="{e(link)}">Open</a>' if str(link).startswith("http") else e(link)
            cells = [f"<b>{e(r['Solicitation Title'])}</b>", e(r["Solicitation Number"]), e(r["Agency"]),
                     e(r["Submission Deadline"]), status_html, e(reason), link_html]
            out.append('<tr valign="top">' + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
        out.append("</table>")
    notes = [(r["Solicitation Title"], r["Manual Review Note"]) for r in rows if str(r.get("Manual Review Note") or "").strip()]
    if notes:
        out.append("<h3>Needs manual review</h3><ul>")
        out += [f"<li><b>{e(t)}</b>: {e(n)}</li>" for t, n in notes]
        out.append("</ul>")
    out.append(f"<p>Full details (all 22 columns) are in <b>{e(tracker_name)}</b>, tab <b>{e(sheet)}</b>, "
               "in the project folder. Inovi facts come only from the four approved Inovi documents.</p>")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(usage=__doc__)
    ap.add_argument("rows")
    ap.add_argument("tracker")
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--html")
    ap.add_argument("--title", default="MyBidMatch Screening Summary")
    a = ap.parse_args()
    sheet = a.sheet.strip()
    if not sheet or len(sheet) > 31 or any(ch in sheet for ch in "[]:*?/\\") or sheet.lower() == "column guide":
        sys.exit("--sheet must be 1-31 characters, not 'Column Guide', without []:*?/\\")
    with open(a.rows, encoding="utf-8") as f:
        rows = json.load(f)
    problems = validate(rows)
    if problems:
        sys.exit("invalid rows:\n" + "\n".join(problems))
    # Render first: a rendering error must not leave the master workbook half-updated.
    page = render_html(rows, a.title, os.path.basename(a.tracker), sheet) if a.html else None
    if page is not None:
        with open(a.html, "w", encoding="utf-8") as f:
            f.write(page)
        print(f"email table written to {a.html}")
    try:
        save_workbook(rows, a.tracker, sheet)
    except WorkbookOpen as e:
        print(f"NOT SAVED: {e}", file=sys.stderr)
        sys.exit(3)
    print(f"tab '{sheet}' with {len(rows)} rows written to {a.tracker}")


if __name__ == "__main__":
    main()
